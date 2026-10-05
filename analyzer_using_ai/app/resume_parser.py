import io
import re
from typing import Dict,List,Optional,Tuple
import unicodedata
from pypdf import PdfReader
from docx import Document


# SECTION_PATTERNS: Dict[str, str] = {
#     "summary": r"\b(summary|profile|professional summary|objective|about me)\b",
#     "education": r"\b(education|academic background|qualification|academic qualifications)\b",
#     "experience": r"\b(experience|work experience|work history|employment|professional experience)\b",
#     "projects": r"\b(projects|project work|personal projects|academic projects)\b",
#     "skills": r"\b(skills|technical skills|core skills|technologies|technical expertise)\b",
#     "certifications": r"\b(certifications|certificates|licenses)\b",
#     "achievements": r"\b(achievements|awards|honors|accomplishments)\b",
#     "contact": r"\b(contact|email|phone|linkedin|github)\b",
# }

# SECTION_ORDER = [
#     "summary",
#     "education",
#     "experience",
#     "projects",
#     "skills",
#     "certifications",
#     "achievements",
# ]

# # find right pattern
# EMAIL_PATTERN = r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}"
# PHONE_PATTERN = r"(\+?\d[\d\s().-]{8,}\d)"

# LINKEDIN_PATTERN = r"(https?://)?(www\.)?linkedin\.com/[A-Za-z0-9_./-]+"

# GITHUB_PATTERN = r"(https?://)?(www\.)?github\.com/[A-Za-z0-9_./-]+"

# PORTFOLIO_PATTERN = r"(https?://[^\s]+)"

# def extract_text_from_pdf(file_bytes:bytes) -> str:
#     reader = PdfReader(io.BytesIO(file_bytes))
#     pages_text = []


#     for page in reader.pages:
#         text = page.extract_text() or ""
#         pages_text.append(text)
    
#     return "\n".join(pages_text)

# # extract text from docx

# def extract_text_from_docx(file_bytes:bytes)-> str:
#     doc = Document(io.BytesIO(file_bytes))
#     #lines = [para.text for para in doc.paragraphs if para.text.strip()]
#     lines = [
#         para.text.strip() for para in doc.paragraphs if para.text.strip()
#     ]
#     return "\n".join(lines)

# # this is function checks whether the file is pdf or docx and extract text from it

# def extract_text(file_name:str,file_bytes:bytes)-> str:
#     lower_name = file_name.lower()

#     if lower_name.endswith(".pdf"):
#         return extract_text_from_pdf(file_bytes)
    
#     if lower_name.endswith(".docx"):
#         return extract_text_from_docx(file_bytes)
    
#     raise ValueError("Unsupported file type. Only PDF and DOCX are allowed.")

# #-----------------------------------
# #text cleaning
# #-----------------------------------

# def clean_text(text: str) -> str:
#     text = text.replace("\x00", " ")

#     #normalize line endingg
#     text = text.replace("\r\n", "\n").replace("\r","\n")
#     #remove excessive spaces
#     text = re.sub(r"[ \t]+", " ", text)
    
#     # text = re.sub(r"[ \t]+", " ", text)
#     # text = re.sub(r"\n{3,}", "\n\n", text)

#     # Remove excessive blank lines
#     text = re.sub(r"\n{3,}", "\n\n", text)

#     return text.strip()


# def clean_line(line:str) -> str:
#     line = line.strip()

#     #remove common bullets
    
#     line = re.sub(r"^[•●▪◦\-–—*]\s*", "", line)

#     return line.strip()


# #--------------------------------------------------------
# # function to detect sections in the resume text
# #---------------------------------------------------------
    



# def detect_sections(text: str) -> List[str]:
#     normalized = text.lower()
#     sections = []

#     for section_name, pattern in SECTION_PATTERNS.items():
#         if re.search(pattern, normalized):
#             sections.append(section_name)

#     return sections


# def is_section_heading(line:str) -> Optional[str]:
#     """
#     Try to determine whether a single line is a section heading.
#     """
#     normalized = line.strip().lower()

#     if not normalized:
#         return None

#     for section_name,pattern in SECTION_PATTERNS.items():
#         if re.fullmatch(pattern,normalized):
#             return section_name

#     # Handle headings such as:
#     # "TECHNICAL SKILLS"
#     # "WORK EXPERIENCE:"

#     normalized_without_colon  = normalized.rstrip(":")

#     for section_name,pattern in SECTION_PATTERNS.items():
#         if re.fullmatch(pattern,normalized_without_colon):
#             return section_name

#     return None


# # splitingg into section

# def split_into_section(text:str) -> Dict[str,List[str]]:
#     """
#     Split resume text into logical sections based on headings.
#     """

#     sections : Dict[str,List[str]] = {}
#     current_section : Optional[str] = None

#     for raw_line in text.split("\n"):
#         line = clean_line(raw_line)

#         if not line:
#             continue

#         detected_section = is_section_heading(line)

#         if detected_section:
#             current_section = detected_section
#             sections.setdefault(current_section,[])
#             continue

#         if current_section:
#             sections.setdefault(current_section,[]).append(line)

#     return sections


# #--------------------------------------------------------------------
# # extract contact info
# #------------------------------------------------------------------

# def extract_email(text:str) -> Optional[str]:
#     match = re.search(EMAIL_PATTERN,text)
#     return match.group(0) if match else None

# def extract_phone(text:str) -> Optional[str]:
#     match = re.search(PHONE_PATTERN,text)

#     if not match:
#         return None
    
#     return re.sub(r"\s+", " ", match.group(0)).strip()



# def has_contact_info(text: str) -> Dict[str, bool]:
#     has_email = bool(re.search(EMAIL_PATTERN, text))
#     has_phone = bool(re.search(PHONE_PATTERN, text))

#     return {
#         "hasEmail": has_email,
#         "hasPhone": has_phone,
#     }


# def extract_link(
#     text: str,
#     pattern: str,
# ) -> Optional[str]:
#     match = re.search(pattern, text, re.IGNORECASE)

#     if not match:
#         return None

#     value = match.group(0).rstrip(".,;)")

#     if not value.startswith("http"):
#         value = f"https://{value}"

#     return value


# def extract_contact_info(text:str) -> Dict[str,Optional[str]]:
#     return {
#         "email":extract_email(text),
#         "phone":extract_phone(text),
#         "linkedln":extract_link(text,LINKEDIN_PATTERN),
#         "github":extract_link(text,GITHUB_PATTERN)
#     }



# #-----------------------------------------------------
# # neme extraction
# #-------------------------------------------------

# def looks_like_name(line:str) -> bool:
#     """
#     Basic heuristic for identifying a candidate's name.

#     This is intentionally conservative. We don't want to
#     accidentally treat email addresses, URLs, headings,
#     or long sentences as a person's name.
#     """

#     line = clean_line(line)

#     if not line:
#         return False

#     if "@" in line:
#         return False

#     if "http://" in line.lower() or "https://" in line.lower():
#         return False

#     if len(line.split()) < 2 or len(line.split()) > 5:
#         return False

#     if len(line) > 60:
#         return False

#     if re.search(r"\d", line):
#         return False

#     section = is_section_heading(line)

#     if section:
#         return False

#     words = line.split()

#     return all(
#         word[0].isupper()
#         for word in words
#         if word and word[0].isalpha()
#     )


# def extract_name(text: str) -> Optional[str]:
#     """
#     Usually the candidate's name appears near the top
#     of the resume.
#     """

#     lines = [
#         clean_line(line)
#         for line in text.split("\n")
#         if clean_line(line)
#     ]

#     # Only inspect the first few lines.
#     for line in lines[:8]:
#         if looks_like_name(line):
#             return line

#     return None



# # ---------------------------------------------------------
# # DATE / PERIOD DETECTION
# # ---------------------------------------------------------

# DATE_RANGE_PATTERN = re.compile(
#     r"""
#     (?P<start>
#         (?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)
#         [a-z]*\s+\d{4}
#         |
#         \d{4}
#     )
#     \s*
#     (?:-|–|—|to)
#     \s*
#     (?P<end>
#         (?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)
#         [a-z]*\s+\d{4}
#         |
#         \d{4}
#         |
#         Present
#         |
#         Current
#     )
#     """,
#     re.IGNORECASE | re.VERBOSE,
# )


# def extract_date_range(text: str) -> Dict[str, Optional[str]]:
#     match = DATE_RANGE_PATTERN.search(text)

#     if not match:
#         return {
#             "startDate": None,
#             "endDate": None,
#         }

#     return {
#         "startDate": match.group("start").strip(),
#         "endDate": match.group("end").strip(),
#     }


# # ---------------------------------------------------------
# # BULLET CLEANING
# # ---------------------------------------------------------

# def clean_bullets(lines: List[str]) -> List[str]:
#     result = []

#     for line in lines:
#         cleaned = clean_line(line)

#         if cleaned:
#             result.append(cleaned)

#     return result


# # ---------------------------------------------------------
# # SUMMARY
# # ---------------------------------------------------------

# def extract_summary(section_lines: List[str]) -> Optional[str]:
#     if not section_lines:
#         return None

#     return " ".join(section_lines).strip()



# # ---------------------------------------------------------
# # SKILLS
# # ---------------------------------------------------------

# def extract_skills(section_lines: List[str]) -> List[str]:
#     """
#     Extract skills from common resume formats.

#     Supports examples like:

#     React, Next.js, Node.js, PostgreSQL

#     or

#     React
#     Next.js
#     Node.js
#     PostgreSQL
#     """

#     skills: List[str] = []

#     separators = r"[,|;/]"

#     for line in section_lines:
#         parts = re.split(separators, line)

#         for part in parts:
#             skill = clean_line(part)

#             if not skill:
#                 continue

#             # Avoid adding extremely long sentences.
#             if len(skill) > 50:
#                 continue

#             if skill.lower() not in {
#                 item.lower()
#                 for item in skills
#             }:
#                 skills.append(skill)

#     return skills

# # ---------------------------------------------------------
# # EDUCATION
# # ---------------------------------------------------------

# def parse_education(section_lines: List[str]) -> List[Dict]:
#     records: List[Dict] = []

#     if not section_lines:
#         return records

#     current: Optional[Dict] = None

#     for line in section_lines:

#         date_info = extract_date_range(line)

#         if date_info["startDate"] or date_info["endDate"]:
#             if current:
#                 records.append(current)

#             remaining = DATE_RANGE_PATTERN.sub("", line).strip(" -|")

#             current = {
#                 "institution": remaining or None,
#                 "degree": None,
#                 "fieldOfStudy": None,
#                 "startDate": date_info["startDate"],
#                 "endDate": date_info["endDate"],
#                 "description": None,
#             }

#             continue

#         # Heuristic: university/college/institute names.
#         lower = line.lower()

#         is_institution = any(
#             keyword in lower
#             for keyword in [
#                 "university",
#                 "college",
#                 "institute",
#                 "school",
#             ]
#         )

#         if current is None:
#             current = {
#                 "institution": line if is_institution else None,
#                 "degree": None,
#                 "fieldOfStudy": None,
#                 "startDate": None,
#                 "endDate": None,
#                 "description": None,
#             }

#             if not is_institution:
#                 current["degree"] = line

#             continue

#         if current["institution"] is None and is_institution:
#             current["institution"] = line
#             continue

#         if current["degree"] is None:
#             current["degree"] = line
#             continue

#         if current["fieldOfStudy"] is None:
#             current["fieldOfStudy"] = line
#             continue

#         if current["description"] is None:
#             current["description"] = line
#         else:
#             current["description"] += f" {line}"

#     if current:
#         records.append(current)

#     return records



# # ---------------------------------------------------------
# # EXPERIENCE
# # ---------------------------------------------------------

# def parse_experience(section_lines: List[str]) -> List[Dict]:
#     records: List[Dict] = []

#     if not section_lines:
#         return records

#     current: Optional[Dict] = None
#     description_lines: List[str] = []

#     for line in section_lines:

#         date_info = extract_date_range(line)

#         if date_info["startDate"] or date_info["endDate"]:

#             if current:
#                 current["description"] = (
#                     " ".join(description_lines).strip() or None
#                 )
#                 records.append(current)

#             before_date = DATE_RANGE_PATTERN.sub("", line).strip(" -|")

#             current = {
#                 "company": None,
#                 "role": before_date or None,
#                 "startDate": date_info["startDate"],
#                 "endDate": date_info["endDate"],
#                 "description": None,
#             }

#             description_lines = []

#             continue

#         if current is None:
#             current = {
#                 "company": None,
#                 "role": line,
#                 "startDate": None,
#                 "endDate": None,
#                 "description": None,
#             }

#             continue

#         # First useful line after role can often be company.
#         if current["company"] is None:
#             current["company"] = line
#             continue

#         description_lines.append(line)

#     if current:
#         current["description"] = (
#             " ".join(description_lines).strip() or None
#         )
#         records.append(current)

#     return records


# # ---------------------------------------------------------
# # PROJECTS
# # ---------------------------------------------------------

# def parse_projects(section_lines: List[str]) -> List[Dict]:
#     records: List[Dict] = []

#     if not section_lines:
#         return records

#     current: Optional[Dict] = None
#     description_lines: List[str] = []

#     for line in section_lines:

#         # A simple heuristic:
#         # lines without bullet-like detail and with short text
#         # are likely project titles.
#         is_possible_title = (
#             len(line) <= 100
#             and not line.endswith(".")
#             and (
#                 ":" in line
#                 or " - " in line
#             )
#         )

#         if current is None:
#             current = {
#                 "title": line,
#                 "description": None,
#                 "technologies": [],
#             }
#             continue

#         if is_possible_title:
#             current["description"] = (
#                 " ".join(description_lines).strip() or None
#             )

#             technologies = extract_technologies(
#                 description_lines
#             )

#             current["technologies"] = technologies

#             records.append(current)

#             current = {
#                 "title": line,
#                 "description": None,
#                 "technologies": [],
#             }

#             description_lines = []

#             continue

#         description_lines.append(line)

#     if current:
#         current["description"] = (
#             " ".join(description_lines).strip() or None
#         )

#         current["technologies"] = extract_technologies(
#             description_lines
#         )

#         records.append(current)

#     return records


# # ---------------------------------------------------------
# # TECHNOLOGY EXTRACTION
# # ---------------------------------------------------------

# KNOWN_TECHNOLOGIES = [
#     "React",
#     "React.js",
#     "Next.js",
#     "Node.js",
#     "Express.js",
#     "TypeScript",
#     "JavaScript",
#     "Python",
#     "FastAPI",
#     "Django",
#     "Flask",
#     "PostgreSQL",
#     "MySQL",
#     "MongoDB",
#     "Prisma",
#     "Docker",
#     "Redis",
#     "AWS",
#     "Git",
#     "GitHub",
#     "Tailwind CSS",
#     "HTML",
#     "CSS",
#     "Redux",
#     "Axios",
#     "TanStack Query",
#     "Socket.IO",
#     "JWT",
#     "Razorpay",
#     "Firebase",
#     "OpenAI",
#     "Gemini",
#     "Ollama",
#     "Pandas",
#     "NumPy",
#     "Scikit-learn",
# ]


# def extract_technologies(lines: List[str]) -> List[str]:
#     text = " ".join(lines)

#     detected = []

#     for technology in KNOWN_TECHNOLOGIES:
#         pattern = rf"(?<!\w){re.escape(technology)}(?!\w)"

#         if re.search(pattern, text, re.IGNORECASE):
#             detected.append(technology)

#     return detected


# # ---------------------------------------------------------
# # CERTIFICATIONS
# # ---------------------------------------------------------

# def parse_certifications(section_lines: List[str]) -> List[Dict]:
#     records = []

#     for line in section_lines:
#         if not line:
#             continue

#         records.append(
#             {
#                 "name": line,
#                 "issuer": None,
#                 "date": None,
#             }
#         )

#     return records


# # ---------------------------------------------------------
# # ACHIEVEMENTS
# # ---------------------------------------------------------

# def parse_achievements(section_lines: List[str]) -> List[str]:
#     return clean_bullets(section_lines)


# # ---------------------------------------------------------
# # STRUCTURED RESUME
# # ---------------------------------------------------------

# def extract_structured_data(
#     text: str,
#     sections: Dict[str, List[str]],
# ) -> Dict:

#     contact = extract_contact_info(text)

#     summary = extract_summary(
#         sections.get("summary", [])
#     )

#     education = parse_education(
#         sections.get("education", [])
#     )

#     experience = parse_experience(
#         sections.get("experience", [])
#     )

#     projects = parse_projects(
#         sections.get("projects", [])
#     )

#     skills = extract_skills(
#         sections.get("skills", [])
#     )

#     certifications = parse_certifications(
#         sections.get("certifications", [])
#     )

#     achievements = parse_achievements(
#         sections.get("achievements", [])
#     )

#     structured_data = {
#         "personal": {
#             "fullName": extract_name(text),
#             "email": contact["email"],
#             "phone": contact["phone"],
#             #"linkedinUrl": contact["linkedin"],
#             "githubUrl": contact["github"],
#             "portfolioUrl": None,
#         },
#         "summary": summary,
#         "education": education,
#         "experience": experience,
#         "projects": projects,
#         "skills": skills,
#         "certifications": certifications,
#         "achievements": achievements,
#     }

#     # Detect an additional portfolio URL.
#     all_urls = re.findall(
#         PORTFOLIO_PATTERN,
#         text,
#         re.IGNORECASE,
#     )

#     for url in all_urls:
#         normalized = url.rstrip(".,;)")

#         if (
#             "linkedin.com" not in normalized.lower()
#             and "github.com" not in normalized.lower()
#         ):
#             structured_data["personal"]["portfolioUrl"] = normalized
#             break

#     return structured_data


# # ---------------------------------------------------------
# # CONTACT FLAGS
# # ---------------------------------------------------------

# def has_contact_info(text: str) -> Dict[str, bool]:
#     has_email = bool(
#         re.search(EMAIL_PATTERN, text)
#     )

#     has_phone = bool(
#         re.search(PHONE_PATTERN, text)
#     )

#     return {
#         "hasEmail": has_email,
#         "hasPhone": has_phone,
#     }


# # ---------------------------------------------------------
# # MAIN PARSER
# # ---------------------------------------------------------

# def parse_resume_file(
#     file_name: str,
#     file_bytes: bytes,
# ) -> Dict:

#     # 1. Extract text
#     raw_text = extract_text(
#         file_name,
#         file_bytes,
#     )

#     # 2. Clean text
#     cleaned_text = clean_text(raw_text)

#     # 3. Basic statistics
#     lines = [
#         line
#         for line in cleaned_text.split("\n")
#         if line.strip()
#     ]

#     words = re.findall(
#         r"\b\w+\b",
#         cleaned_text,
#     )

#     # 4. Detect sections
#     sections_found = detect_sections(
#         cleaned_text
#     )

#     # 5. Split sections
#     sections = split_into_section(
#         cleaned_text
#     )

#     # 6. Contact information
#     contact_info = has_contact_info(
#         cleaned_text
#     )

#     # 7. Structured information
#     structured_data = extract_structured_data(
#         cleaned_text,
#         sections,
#     )

#     return {
#         "fileName": file_name,
#         "fileType": file_name.split(".")[-1].lower(),
#         "rawText": raw_text,
#         "cleanedText": cleaned_text,
#         "wordCount": len(words),
#         "lineCount": len(lines),
#         "sectionsFound": sections_found,
#         "sectionData": sections,
#         "structuredData": structured_data,
#         **contact_info,
#     }




# =========================================================
# SECTION DEFINITIONS
# =========================================================

SECTION_ALIASES = {
    "summary": [
        "summary",
        "profile",
        "professional summary",
        "objective",
        "career objective",
        "about me",
    ],
    "education": [
        "education",
        "academic background",
        "qualification",
        "academic qualifications",
    ],
    "experience": [
        "experience",
        "work experience",
        "work history",
        "employment",
        "professional experience",
    ],
    "projects": [
        "projects",
        "project work",
        "personal projects",
        "academic projects",
    ],
    "skills": [
        "skills",
        "technical skills",
        "core skills",
        "technologies",
        "technical expertise",
    ],
    "certifications": [
        "certifications",
        "certificates",
        "licenses",
    ],
    "achievements": [
        "achievements",
        "awards",
        "honors",
        "accomplishments",
    ],
}


SECTION_HEADING_MAP = {
    alias.lower(): section
    for section, aliases in SECTION_ALIASES.items()
    for alias in aliases
}


SECTION_PATTERN = re.compile(
    r"^("
    + "|".join(
        re.escape(alias)
        for alias in sorted(
            SECTION_HEADING_MAP.keys(),
            key=len,
            reverse=True,
        )
    )
    + r")\s*:?\s*$",
    re.IGNORECASE,
)


# =========================================================
# REGEX PATTERNS
# =========================================================

EMAIL_PATTERN = (
    r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}"
)

PHONE_PATTERN = (
    r"(?:\+\d{1,3}[\s-]?)?"
    r"(?:\(?\d{2,4}\)?[\s.-]?)?"
    r"\d{3,4}[\s.-]?\d{3,4}"
)


DATE_RANGE_PATTERN = re.compile(
    r"""
    (?P<start>
        (?:
            Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec
        )[a-z]*\s+\d{4}
        |
        \d{4}
    )
    \s*
    (?:-|–|—|to)
    \s*
    (?P<end>
        (?:
            Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec
        )[a-z]*\s+\d{4}
        |
        \d{4}
        |
        Present
        |
        Current
    )
    """,
    re.IGNORECASE | re.VERBOSE,
)


LOCATION_PATTERN = re.compile(
    r"(?P<location>"
    r"[A-Za-z][A-Za-z .'-]{1,40},\s*"
    r"[A-Za-z][A-Za-z .'-]{1,40}"
    r")\s*$"
)


INSTITUTION_KEYWORDS = (
    "university",
    "college",
    "institute",
    "school",
)


# =========================================================
# SKILL NORMALIZATION
# =========================================================

CANONICAL_SKILLS = {
    "javascript": "JavaScript",
    "angular": "Angular",
    "nodejs": "Node.js",
    "node.js": "Node.js",
    "mysql": "MySQL",
    "api gateway": "API Gateway",
    "git": "Git",
    "python": "Python",
    "docker": "Docker",
    "sonarqube": "SonarQube",
    "model check": "Model Check",
    "typescript": "TypeScript",
    "react": "React",
    "react.js": "React",
    "next.js": "Next.js",
    "postgresql": "PostgreSQL",
    "mongodb": "MongoDB",
    "aws": "AWS",
    "redis": "Redis",
    "php": "PHP",
    "fastapi": "FastAPI",
    "sql": "SQL",
    "html": "HTML",
    "css": "CSS",
    "redux": "Redux",
}


# =========================================================
# TEXT EXTRACTION
# =========================================================

def normalize_unicode(text: str) -> str:
    """
    Convert ligatures and other compatible Unicode characters
    into normal text.

    Example:
        Certiﬁed -> Certified
        trafﬁc   -> traffic
    """

    return unicodedata.normalize("NFKC", text or "")


def extract_text_from_pdf(file_bytes: bytes) -> str:
    reader = PdfReader(io.BytesIO(file_bytes))

    pages_text = []

    for page in reader.pages:
        text = page.extract_text() or ""
        pages_text.append(text)

    return "\n".join(pages_text)


def extract_text_from_docx(file_bytes: bytes) -> str:
    document = Document(io.BytesIO(file_bytes))

    lines = [
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]

    return "\n".join(lines)


def extract_text(
    file_name: str,
    file_bytes: bytes,
) -> str:

    lower_name = file_name.lower()

    if lower_name.endswith(".pdf"):
        return extract_text_from_pdf(file_bytes)

    if lower_name.endswith(".docx"):
        return extract_text_from_docx(file_bytes)

    raise ValueError(
        "Unsupported file type. Only PDF and DOCX are allowed."
    )


# =========================================================
# TEXT CLEANING
# =========================================================

def clean_text(text: str) -> str:

    text = normalize_unicode(text)

    text = text.replace("\x00", " ")

    text = text.replace("\r\n", "\n")
    text = text.replace("\r", "\n")

    text = re.sub(
        r"[ \t]+",
        " ",
        text,
    )

    text = re.sub(
        r"\n{3,}",
        "\n\n",
        text,
    )

    return text.strip()


def clean_line(line: str) -> str:

    line = line.strip()

    # Remove common bullet characters.
    line = re.sub(
        r"^[•●▪◦\-–—*]\s*",
        "",
        line,
    )

    return line.strip()


# =========================================================
# SECTION DETECTION
# =========================================================

def is_section_heading(
    line: str,
) -> Optional[str]:

    cleaned = clean_line(line)

    match = SECTION_PATTERN.match(cleaned)

    if not match:
        return None

    return SECTION_HEADING_MAP[
        match.group(1).lower()
    ]


def split_into_sections(
    text: str,
) -> Tuple[Dict[str, List[str]], List[str]]:

    sections: Dict[str, List[str]] = {}

    current_section: Optional[str] = None

    for raw_line in text.split("\n"):

        line = clean_line(raw_line)

        if not line:
            continue

        detected_section = is_section_heading(line)

        if detected_section:

            current_section = detected_section

            sections.setdefault(
                current_section,
                [],
            )

            continue

        if current_section:
            sections.setdefault(
                current_section,
                [],
            ).append(line)

    return sections, list(sections.keys())


# =========================================================
# CONTACT INFORMATION
# =========================================================

def extract_email(
    text: str,
) -> Optional[str]:

    match = re.search(
        EMAIL_PATTERN,
        text,
    )

    return match.group(0) if match else None


def extract_phone(
    text: str,
) -> Optional[str]:

    match = re.search(
        PHONE_PATTERN,
        text,
    )

    if not match:
        return None

    return match.group(0).strip()


def extract_url(
    text: str,
    host: str,
) -> Optional[str]:

    pattern = (
        rf"(?:https?://)?"
        rf"(?:www\.)?"
        rf"{re.escape(host)}"
        rf"/[A-Za-z0-9_./-]+"
    )

    match = re.search(
        pattern,
        text,
        re.IGNORECASE,
    )

    if not match:
        return None

    value = match.group(0).rstrip(
        ".,;)"
    )

    if not value.startswith("http"):
        value = f"https://{value}"

    return value


def extract_location(
    lines: List[str],
) -> Optional[str]:

    for line in lines[:12]:

        if re.fullmatch(
            r"[A-Za-z][A-Za-z .'-]{1,30},\s*"
            r"[A-Za-z][A-Za-z .'-]{1,30}",
            line,
        ):
            return line

    return None


# =========================================================
# NAME EXTRACTION
# =========================================================

def extract_name(
    lines: List[str],
) -> Optional[str]:

    for line in lines[:6]:

        line = clean_line(line)

        if not line:
            continue

        if "@" in line:
            continue

        if re.search(r"\d", line):
            continue

        if is_section_heading(line):
            continue

        words = line.split()

        if not (
            2 <= len(words) <= 5
        ):
            continue

        if len(line) > 60:
            continue

        if all(
            word[0].isupper()
            for word in words
            if word and word[0].isalpha()
        ):
            return line

    return None


# =========================================================
# DATE + LOCATION
# =========================================================

def split_date_and_location(
    line: str,
):

    match = DATE_RANGE_PATTERN.search(line)

    if not match:
        return None

    start = match.group("start").strip()
    end = match.group("end").strip()

    remaining = line[
        match.end():
    ].strip(" |,")

    location = None

    if remaining:
        location_match = LOCATION_PATTERN.search(
            remaining
        )

        if location_match:
            location = location_match.group(
                "location"
            )

    return start, end, location


# =========================================================
# WRAPPED LINE HANDLING
# =========================================================

def merge_wrapped_lines(
    lines: List[str],
) -> List[str]:

    result: List[str] = []

    for raw_line in lines:

        line = clean_line(raw_line)

        if not line:
            continue

        if result:

            # PDF extraction often breaks a sentence
            # across multiple lines. When the new line
            # starts with a lowercase character, treat it
            # as a continuation.
            if line[0].islower():

                result[-1] = (
                    result[-1]
                    + " "
                    + line
                )

                continue

        result.append(line)

    return result


# =========================================================
# SUMMARY
# =========================================================

def extract_summary(
    lines: List[str],
) -> Optional[str]:

    merged = merge_wrapped_lines(lines)

    if not merged:
        return None

    return " ".join(merged).strip()


# =========================================================
# EDUCATION
# =========================================================

# def parse_education(
#     lines: List[str],
# ) -> List[Dict]:

#     if not lines:
#         return []

#     lines = [
#         clean_line(line)
#         for line in lines
#         if clean_line(line)
#     ]

#     if not lines:
#         return []

#     institution = None
#     degree = None
#     field = None
#     start_date = None
#     end_date = None
#     location = None

#     date_index = next(
#         (
#             index
#             for index, line in enumerate(lines)
#             if DATE_RANGE_PATTERN.search(line)
#         ),
#         None,
#     )

#     if date_index is not None:

#         date_info = split_date_and_location(
#             lines[date_index]
#         )

#         if date_info:
#             start_date = date_info[0]
#             end_date = date_info[1]
#             location = date_info[2]

#     before_date = (
#         lines[:date_index]
#         if date_index is not None
#         else lines
#     )

#     after_date = (
#         lines[date_index + 1 :]
#         if date_index is not None
#         else []
#     )

#     before_date = merge_wrapped_lines(
#         before_date
#     )

#     institution_index = next(
#         (
#             index
#             for index, line in enumerate(before_date)
#             if any(
#                 keyword in line.lower()
#                 for keyword in INSTITUTION_KEYWORDS
#             )
#         ),
#         None,
#     )

#     institution = None
#     degree = None
#     field_of_study = None

#     if institution_index is not None:

#         institution = before_date[
#             institution_index
#         ]

#         # Join a wrapped institution name.
#         if (
#             institution_index + 1
#             < len(before_date)
#             and len(
#                 before_date[
#                     institution_index + 1
#                 ].split()
#             ) <= 5
#         ):
#             next_line = before_date[
#                 institution_index + 1
#             ]

#             institution = (
#                 institution
#                 + " "
#                 + next_line
#             )

#         degree = (
#             before_date[0]
#             if institution_index >= 1
#             else None
#         )

#         field = " ".join(
#             before_date[
#                 1:institution_index
#             ]
#         ).strip() or None

#     else:

#         if before_date:
#             degree = before_date[0]

#         if len(before_date) > 1:
#             field = " ".join(
#                 before_date[1:]
#             )

#         if after_date:
#             institution = after_date[0]

#     start_date = None
#     end_date = None
#     location = None

#     if date_index is not None:

#         date_info = split_date_and_location(
#             lines[date_index]
#         )

#         if date_info:

#             start_date = date_info[0]
#             end_date = date_info[1]
#             location = date_info[2]

#     # Some documents put the location on the
#     # line after the date.
#     if location is None and after_date:

#         possible_location = after_date[0]

#         if re.fullmatch(
#             r"[A-Za-z][A-Za-z .'-]{1,30},\s*"
#             r"[A-Za-z][A-Za-z .'-]{1,30}",
#             possible_location,
#         ):
#             location = possible_location

#     return [
#         {
#             "institution": institution,
#             "degree": degree,
#             "fieldOfStudy": field,
#             "startDate": start_date,
#             "endDate": end_date,
#             "location": location,
#         }
#     ]

def parse_education(
    lines: List[str],
) -> List[Dict]:

    if not lines:
        return []

    lines = [
        clean_line(line)
        for line in lines
        if clean_line(line)
    ]

    if not lines:
        return []

    institution = None
    degree = None
    field = None
    start_date = None
    end_date = None
    location = None

    # -----------------------------------------------------
    # Find the line containing the education date range
    # -----------------------------------------------------

    date_index = next(
        (
            index
            for index, line in enumerate(lines)
            if DATE_RANGE_PATTERN.search(line)
        ),
        None,
    )

    # -----------------------------------------------------
    # Extract dates
    # -----------------------------------------------------

    if date_index is not None:

        date_info = split_date_and_location(
            lines[date_index]
        )

        if date_info:
            start_date = date_info[0]
            end_date = date_info[1]
            location = date_info[2]

    # -----------------------------------------------------
    # Remove date/location line from education content
    # -----------------------------------------------------

    content_lines = []

    for index, line in enumerate(lines):

        if index == date_index:
            continue

        content_lines.append(line)

    content_lines = merge_wrapped_lines(
        content_lines
    )

    if not content_lines:
        return []

    # -----------------------------------------------------
    # Find institution
    # -----------------------------------------------------

    institution_index = next(
        (
            index
            for index, line in enumerate(content_lines)
            if any(
                keyword in line.lower()
                for keyword in INSTITUTION_KEYWORDS
            )
        ),
        None,
    )

    # -----------------------------------------------------
    # Institution found
    # -----------------------------------------------------

    if institution_index is not None:

        institution = content_lines[
            institution_index
        ]

        # If institution wraps onto the next line,
        # join it with the current line.
        if (
            institution_index + 1 < len(content_lines)
            and institution_index + 1 != len(content_lines)
        ):

            next_line = content_lines[
                institution_index + 1
            ]

            # Only join a short continuation line.
            if (
                len(next_line.split()) <= 5
                and not any(
                    keyword in next_line.lower()
                    for keyword in (
                        "bachelor",
                        "master",
                        "phd",
                        "doctor",
                        "diploma",
                        "certificate",
                    )
                )
            ):
                institution = (
                    institution
                    + " "
                    + next_line
                )

        # -------------------------------------------------
        # Everything before institution
        # -------------------------------------------------

        before_institution = content_lines[
            :institution_index
        ]

        if before_institution:

            degree = before_institution[0]

            if len(before_institution) > 1:

                field = " ".join(
                    before_institution[1:]
                ).strip()

        # -------------------------------------------------
        # If no clear field was found
        # -------------------------------------------------

        if field == "":
            field = None

    # -----------------------------------------------------
    # Institution not detected
    # -----------------------------------------------------

    else:

        if len(content_lines) >= 1:
            degree = content_lines[0]

        if len(content_lines) >= 2:
            field = content_lines[1]

        if len(content_lines) >= 3:
            institution = content_lines[2]

    # -----------------------------------------------------
    # Extract location from remaining lines
    # -----------------------------------------------------

    if location is None:

        for line in content_lines:

            if re.fullmatch(
                r"[A-Za-z][A-Za-z .'-]{1,30},\s*"
                r"[A-Za-z][A-Za-z .'-]{1,30}",
                line,
            ):
                location = line
                break

    # -----------------------------------------------------
    # Return one education record
    # -----------------------------------------------------

    return [
        {
            "institution": institution,
            "degree": degree,
            "fieldOfStudy": field,
            "startDate": start_date,
            "endDate": end_date,
            "location": location,
        }
    ]


# =========================================================
# EXPERIENCE
# =========================================================

def parse_experience(
    lines: List[str],
) -> List[Dict]:

    if not lines:
        return []

    lines = [
        clean_line(line)
        for line in lines
        if clean_line(line)
    ]

    date_positions = [
        index
        for index, line in enumerate(lines)
        if DATE_RANGE_PATTERN.search(line)
    ]

    records: List[Dict] = []

    for position, date_index in enumerate(
        date_positions
    ):

        # Job title and company are normally the
        # two lines immediately before the date.
        header_start = max(
            0,
            date_index - 2,
        )

        header_lines = lines[
            header_start:date_index
        ]

        role = (
            header_lines[-2]
            if len(header_lines) >= 2
            else (
                header_lines[-1]
                if header_lines
                else None
            )
        )

        company = (
            header_lines[-1]
            if len(header_lines) >= 2
            else None
        )

        date_info = split_date_and_location(
            lines[date_index]
        )

        if date_info:

            start_date = date_info[0]
            end_date = date_info[1]
            location = date_info[2]

        else:

            start_date = None
            end_date = None
            location = None

        # The next two lines before the next date
        # belong to the next role/company pair.
        if position + 1 < len(date_positions):

            next_date_index = date_positions[
                position + 1
            ]

            responsibility_end = max(
                date_index + 1,
                next_date_index - 2,
            )

        else:

            responsibility_end = len(lines)

        responsibility_lines = lines[
            date_index + 1:
            responsibility_end
        ]

        responsibilities = merge_wrapped_lines(
            responsibility_lines
        )

        records.append(
            {
                "company": company,
                "role": role,
                "startDate": start_date,
                "endDate": end_date,
                "location": location,
                "responsibilities": responsibilities,
            }
        )

    return records


# =========================================================
# SKILLS
# =========================================================

def extract_skills(
    lines: List[str],
) -> List[str]:

    found: List[str] = []

    for line in lines:

        parts = re.split(
            r"[,|;/]",
            line,
        )

        for part in parts:

            skill = clean_line(part)

            if not skill:
                continue

            key = skill.lower()

            canonical = CANONICAL_SKILLS.get(
                key,
                skill,
            )

            # Do not treat long sentences as skills.
            if len(canonical) > 50:
                continue

            if canonical.lower() not in {
                item.lower()
                for item in found
            }:
                found.append(canonical)

    return found


# =========================================================
# CERTIFICATIONS
# =========================================================

def parse_certifications(
    lines: List[str],
) -> List[Dict]:

    records = []

    for line in lines:

        title = clean_line(line)

        if not title:
            continue

        records.append(
            {
                "title": title,
                "issuer": None,
                "date": None,
            }
        )

    return records


# =========================================================
# PROJECTS
# =========================================================

def parse_projects(
    lines: List[str],
) -> List[Dict]:

    # The first version intentionally returns an empty
    # list when there is no dedicated project parser yet.
    #
    # We will improve project extraction after validating
    # education and experience across multiple resumes.

    return []


# =========================================================
# ACHIEVEMENTS
# =========================================================

def parse_achievements(
    lines: List[str],
) -> List[str]:

    return [
        clean_line(line)
        for line in lines
        if clean_line(line)
    ]


# =========================================================
# STRUCTURED DATA
# =========================================================

def extract_structured_data(
    text: str,
    sections: Dict[str, List[str]],
    top_lines: List[str],
) -> Dict:

    email = extract_email(text)
    phone = extract_phone(text)

    personal = {
        "fullName": extract_name(top_lines),
        "email": email,
        "phone": phone,
        "linkedinUrl": extract_url(
            text,
            "linkedin.com",
        ),
        "githubUrl": extract_url(
            text,
            "github.com",
        ),
        "portfolioUrl": None,
        "location": extract_location(
            top_lines
        ),
    }

    # Only use the dedicated Skills section
    # for profile population.
    skills = extract_skills(
        sections.get("skills", [])
    )

    return {
        "personal": personal,

        "summary": extract_summary(
            sections.get("summary", [])
        ),

        "education": parse_education(
            sections.get("education", [])
        ),

        "experience": parse_experience(
            sections.get("experience", [])
        ),

        "projects": parse_projects(
            sections.get("projects", [])
        ),

        "skills": skills,

        "certifications": parse_certifications(
            sections.get("certifications", [])
        ),

        "achievements": parse_achievements(
            sections.get("achievements", [])
        ),
    }


# =========================================================
# CONTACT FLAGS
# =========================================================

def has_contact_info(
    text: str,
) -> Dict[str, bool]:

    return {
        "hasEmail": bool(
            re.search(
                EMAIL_PATTERN,
                text,
            )
        ),
        "hasPhone": bool(
            re.search(
                PHONE_PATTERN,
                text,
            )
        ),
    }


# =========================================================
# MAIN PARSER
# =========================================================

def parse_resume_file(
    file_name: str,
    file_bytes: bytes,
) -> Dict:

    # Step 1: Extract text
    raw_text = extract_text(
        file_name,
        file_bytes,
    )

    # Step 2: Clean extracted text
    cleaned_text = clean_text(
        raw_text
    )

    # Step 3: Get individual lines
    all_lines = [
        clean_line(line)
        for line in cleaned_text.split("\n")
        if clean_line(line)
    ]

    # Step 4: Detect and split sections
    sections, sections_found = (
        split_into_sections(
            cleaned_text
        )
    )

    # Step 5: Structured information
    structured_data = extract_structured_data(
        cleaned_text,
        sections,
        all_lines,
    )

    # Step 6: Contact information flags
    contact_info = has_contact_info(
        cleaned_text
    )

    # Step 7: Return complete parser response
    return {
        "fileName": file_name,

        "fileType": file_name
        .split(".")[-1]
        .lower(),

        "rawText": raw_text,

        "cleanedText": cleaned_text,

        "wordCount": len(
            re.findall(
                r"\b\w+\b",
                cleaned_text,
            )
        ),

        "lineCount": len(all_lines),

        "sectionsFound": sections_found,

        "sectionData": sections,

        "structuredData": structured_data,

        **contact_info,
    }